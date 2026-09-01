"""
reports/audit_export.py
Server-side Excel (.xlsx) and CSV export for the Audit History feature —
one row per audited session, built from the SAME filters as the history
list and dashboard (history_filters.HistoryFilters), fetched in batches via
session_store.iter_sessions_for_export() so the Postgres side stays flat
regardless of how many rows match (60 today, architected for 10,000+): the
query is paginated internally (same keyset mechanism as "Load more").

The XLSX has three sheets:
  1. "Audit Results"  — the manager-facing summary (EXPORT_COLUMNS, below).
  2. "Tracker Format"  — the exact Audit Tracker / Google Sheet column
     schema (reports/audit_excel_manager.COLUMNS), one row per session,
     built via that module's own auto_fill_from_report() — the SAME
     function the per-audit "Push to Tracker" feature already uses, so
     this sheet is guaranteed to match what actually gets pushed, not a
     second reimplementation. Requires each session's full report_json
     (deck_coverage/duration/talk-ratio/etc, not just the 3 summary
     fields) — see iter_sessions_for_export(include_full_report=True).
     Video Snapshot Analysis fields (Demo Attendees, Camera Status,
     Screen-share Mode) are left blank here even when that add-on ran for
     a session — auto_fill_from_report() accepts that data as an optional
     argument the export doesn't currently look up per session (would be
     an extra query per row); everything else on this sheet is complete.
  3. "Export Summary"  — filters used + cheap aggregate metrics.
CSV export has none of this — one flat table only (EXPORT_COLUMNS).

Column order and content for the main sheet follow the feature spec
exactly — see CATEGORY_EXPORT_ORDER, derived directly from
audit.evaluation_framework.EVALUATION_CATEGORIES (not a hand-copied
parallel list) so the two can never drift apart.

Privacy: deliberately excludes PII the CRM import carries but the spec
doesn't ask for — Email, Phone Number, Mobile Number are NEVER exported
here even though they're present in lead_sheet_json for CSV-uploaded
sessions. This export is manager-facing (scores/performance/CRM
identifiers), not a raw dump of the CRM row.

"Associate Email" maps from lead_sheet_json's "Owner Email" — the CSV has
no separate "Lead Owner Email" field, and real production data confirmed
Owner and Lead Owner are always the same person, so Owner Email is the
correct (only available) email for the canonical associate.
"""
from __future__ import annotations
import csv
import io
from datetime import datetime, timezone

import session_store
from audit.evaluation_framework import EVALUATION_CATEGORIES

CATEGORY_EXPORT_ORDER = [(c["id"], c["label"]) for c in EVALUATION_CATEGORIES]

EXPORT_COLUMNS = (
    ["Audit Date", "Demo Date", "Associate", "Associate Email", "TL",
     "Lead Name", "Prospect ID", "Lead Number", "Lead Stage", "Payment Done",
     "Session ID", "Audit Score", "Grade", "Strongest Category", "Weakest Category"]
    + [label for _, label in CATEGORY_EXPORT_ORDER]
    + ["Strengths", "Areas to Improve", "Lead Link", "Demo Link"]
)

EXPORT_BATCH_SIZE = 500  # rows fetched from Postgres per round trip while exporting


def _fmt_ts(value: str) -> str:
    """ISO-ish stored timestamp -> 'DD Mon YYYY HH:MM' for the export, or
    '' if blank/unparseable — never raises on a slightly different format
    (e.g. lead_sheet's space-separated CRM timestamps vs the app's own
    'T'-separated ones)."""
    if not value:
        return ""
    v = value.replace("T", " ")
    for fmt in ("%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(v, fmt).strftime("%d %b %Y %H:%M")
        except ValueError:
            continue
    return value  # fall back to the raw stored value rather than hiding it


def _row_for_session(rec: dict) -> dict:
    """Builds one export row (dict keyed by EXPORT_COLUMNS) from a
    session_store.iter_sessions_for_export() record — no further DB calls,
    everything needed is already on `rec` (lightweight fields +
    category_scores_json/top_strengths_json/improvement_areas_json +
    lead_sheet, all fetched in bulk upstream)."""
    lead = rec.get("lead_sheet") or {}
    cat_scores = rec.get("category_scores_json") or {}

    def cat_score(cat_id):
        c = cat_scores.get(cat_id)
        if isinstance(c, dict):
            v = c.get("raw_score")
            return round(v, 1) if isinstance(v, (int, float)) else ""
        return ""

    ranked = sorted(
        ((label, cat_scores[cid]["raw_score"]) for cid, label in CATEGORY_EXPORT_ORDER
         if isinstance(cat_scores.get(cid), dict) and isinstance(cat_scores[cid].get("raw_score"), (int, float))),
        key=lambda kv: -kv[1],
    )
    strongest = ranked[0][0] if ranked else ""
    weakest = ranked[-1][0] if ranked else ""

    strengths = rec.get("top_strengths_json") or []
    improvements = rec.get("improvement_areas_json") or []

    row = {
        "Audit Date":        _fmt_ts(rec.get("completed_at") or rec.get("created_at") or ""),
        "Demo Date":         _fmt_ts(lead.get("Demo Date", "")),
        "Associate":         rec.get("label") or "",
        "Associate Email":   lead.get("Owner Email", ""),
        "TL":                lead.get("TL Name", ""),
        "Lead Name":         lead.get("Lead Name", ""),
        "Prospect ID":       lead.get("Prospect ID", ""),
        "Lead Number":       lead.get("Lead Number", ""),
        "Lead Stage":        lead.get("Lead Stage", ""),
        "Payment Done":      lead.get("Payment Done", ""),
        "Session ID":        rec.get("session_id", ""),
        "Audit Score":       rec.get("overall_score") if rec.get("overall_score") is not None else "",
        "Grade":             rec.get("grade") or "",
        "Strongest Category": strongest,
        "Weakest Category":   weakest,
        "Strengths":         " | ".join(strengths) if strengths else "",
        "Areas to Improve":  " | ".join(improvements) if improvements else "",
        "Lead Link":         lead.get("Lead link", ""),
        "Demo Link":         rec.get("source_url") or "",
    }
    for cid, label in CATEGORY_EXPORT_ORDER:
        row[label] = cat_score(cid)
    return row


def _summary_rows(filters, total: int) -> list[tuple]:
    """(label, value) rows for the Summary sheet — reuses the dashboard's
    own aggregation (reports/associate_analytics.build_overall_dashboard,
    same filters) rather than recomputing average/strongest/weakest a
    second, different way."""
    from reports.associate_analytics import build_overall_dashboard
    dash = build_overall_dashboard(filters) if total else {}
    grade_dist = dash.get("grade_distribution") or {}
    most_common_grade = max(grade_dist, key=grade_dist.get) if grade_dist else "—"

    lines: list[tuple] = [
        ("Kalvium Audit Report", ""),
        ("Exported At", datetime.now(timezone.utc).strftime("%d %b %Y %H:%M UTC")),
        ("Total Matching Audits", total),
        ("", ""),
        ("Filters", ""),
    ]
    for label, value in (filters.as_display_lines() if filters else []):
        lines.append((label, value))
    lines += [
        ("", ""),
        ("Average Score",     dash.get("average_score", "—")),
        ("Total Audits",      dash.get("total_audits", total)),
        ("Most Common Grade", most_common_grade),
        ("Strongest Category", (dash.get("strongest_categories") or ["—"])[0]),
        ("Weakest Category",   (dash.get("weakest_categories") or ["—"])[0]),
    ]
    return lines


def _export_filename(fmt: str, filters=None) -> str:
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    hint = ""
    if filters and not filters.is_empty():
        if len(filters.associate) == 1:
            safe = "".join(c for c in filters.associate[0] if c.isalnum() or c in " _-").strip().replace(" ", "_")[:24]
            if safe:
                hint = f"_{safe}"
        elif filters.associate or filters.tl or filters.lead_stage or filters.payment_status or filters.search:
            hint = "_filtered"
    return f"kalvium_audit_export_{today}{hint}.{fmt}"


def generate_csv(filters=None) -> tuple[bytes, str]:
    """Returns (utf-8 CSV bytes, filename). No summary — CSV is one table only."""
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=EXPORT_COLUMNS, extrasaction="ignore")
    writer.writeheader()
    for batch in session_store.iter_sessions_for_export(filters, batch_size=EXPORT_BATCH_SIZE):
        for rec in batch:
            writer.writerow(_row_for_session(rec))
    return buf.getvalue().encode("utf-8-sig"), _export_filename("csv", filters)  # BOM: Excel opens UTF-8 CSVs correctly


def _tracker_row_for_session(rec: dict) -> dict:
    """Builds one row matching reports.audit_excel_manager.COLUMNS exactly
    — reuses that module's own auto_fill_from_report(), the SAME logic the
    per-audit "Push to Tracker" feature already uses, so the Tracker Format
    sheet can never drift from what a real tracker push actually produces."""
    from reports.audit_excel_manager import auto_fill_from_report
    report = dict(rec.get("full_report") or {})
    report.setdefault("session_id", rec.get("session_id", ""))
    session_meta = {
        "source_url": rec.get("source_url") or "",
        "label": rec.get("label") or "",
        "lead_sheet": rec.get("lead_sheet") or {},
        "video_analysis": None,  # see module docstring: not looked up per export row
    }
    return auto_fill_from_report(report, session_meta)


def generate_xlsx(filters=None) -> tuple[bytes, str]:
    """Returns (xlsx bytes, filename). Three sheets: 'Audit Results' (one
    row per session, frozen header, autofilter, sized columns, wrapped
    long-text columns, clickable Lead/Demo Link), 'Tracker Format' (the
    exact Audit Tracker column schema — see _tracker_row_for_session), and
    'Export Summary' (filters used + cheap aggregate metrics, reusing the
    dashboard's own numbers).

    Normal (non-write-only) workbook mode — write-only mode was tried
    first for the memory-flatness Section 21 asks for, but openpyxl's
    write-only worksheets silently drop freeze_panes on save (confirmed:
    reloading a write-only-saved file always comes back with
    freeze_panes=None, a real library limitation, not a usage mistake) —
    and freeze_panes is an explicit requirement here. The batched Postgres
    fetch (session_store.iter_sessions_for_export) still keeps the
    DATABASE-side memory flat regardless of row count; only the final
    in-memory workbook object scales with rows, which is normal for any
    XLSX library and not a practical concern at the 10K-row target scale
    (tens of MB, not hundreds)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    from reports.audit_excel_manager import COLUMNS as TRACKER_COLUMNS, _header_bg

    total = session_store.count_sessions_matching(filters)

    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Results"

    header_font = Font(bold=True)
    for i, col_name in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=col_name)
        cell.font = header_font

    wrap_cols = {"Strengths", "Areas to Improve", "Lead Name"}
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    link_font = Font(color="0563C1", underline="single")
    col_widths = {
        "Audit Date": 17, "Demo Date": 17, "Associate": 22, "Associate Email": 24,
        "TL": 18, "Lead Name": 26, "Prospect ID": 24, "Lead Number": 14,
        "Lead Stage": 16, "Payment Done": 13, "Session ID": 24, "Audit Score": 11,
        "Grade": 8, "Strongest Category": 20, "Weakest Category": 20,
        "Strengths": 45, "Areas to Improve": 45, "Lead Link": 30, "Demo Link": 30,
    }
    for i, col_name in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = col_widths.get(col_name, 13)

    # Tracker Format sheet — same styling approach (header background per
    # column) as the live master tracker (reports/audit_excel_manager),
    # via that module's own _header_bg() so the two never look different.
    tracker_ws = wb.create_sheet("Tracker Format")
    tracker_wrap = Alignment(wrap_text=True, vertical="top")
    for i, col_name in enumerate(TRACKER_COLUMNS, start=1):
        bg, fg = _header_bg(i)
        cell = tracker_ws.cell(row=1, column=i, value=col_name.replace("\n", " ").strip())
        cell.font = Font(bold=True, color=fg)
        cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        cell.alignment = tracker_wrap
        tracker_ws.column_dimensions[get_column_letter(i)].width = 16

    row_idx = 1
    for batch in session_store.iter_sessions_for_export(filters, batch_size=EXPORT_BATCH_SIZE,
                                                          include_full_report=True):
        for rec in batch:
            row_idx += 1
            row = _row_for_session(rec)
            for i, col_name in enumerate(EXPORT_COLUMNS, start=1):
                value = row.get(col_name, "")
                cell = ws.cell(row=row_idx, column=i, value=value)
                if col_name in wrap_cols:
                    cell.alignment = wrap_alignment
                if col_name in ("Lead Link", "Demo Link") and value:
                    cell.hyperlink = value
                    cell.font = link_font

            tracker_row = _tracker_row_for_session(rec)
            for i, col_name in enumerate(TRACKER_COLUMNS, start=1):
                tracker_ws.cell(row=row_idx, column=i, value=tracker_row.get(col_name, ""))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}{max(row_idx, 1)}"
    tracker_ws.freeze_panes = "A2"
    tracker_ws.auto_filter.ref = f"A1:{get_column_letter(len(TRACKER_COLUMNS))}{max(row_idx, 1)}"

    summary_ws = wb.create_sheet("Export Summary")
    for r, (label, value) in enumerate(_summary_rows(filters, total), start=1):
        label_cell = summary_ws.cell(row=r, column=1, value=label)
        summary_ws.cell(row=r, column=2, value=value)
        if label == "Kalvium Audit Report":
            label_cell.font = Font(bold=True, size=14)
        elif label and not value:
            label_cell.font = Font(bold=True)  # section headers ("Filters", blank-value rows)
    summary_ws.column_dimensions["A"].width = 22
    summary_ws.column_dimensions["B"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), _export_filename("xlsx", filters)
