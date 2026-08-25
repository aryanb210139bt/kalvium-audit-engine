"""
deck/ingestion/pptx_extractor.py
Extracts raw content from a .pptx file slide by slide.
"""
from __future__ import annotations
import logging
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)


@dataclass
class RawSlide:
    slide_number: int
    title: str = ""
    body_text: str = ""
    speaker_notes: str = ""
    table_cells: list[str] = field(default_factory=list)
    image_alts: list[str] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        parts = []
        if self.title:
            parts.append(f"[Title] {self.title}")
        if self.body_text:
            parts.append(self.body_text)
        if self.table_cells:
            parts.append(" | ".join(self.table_cells))
        if self.speaker_notes:
            parts.append(f"[Notes] {self.speaker_notes}")
        return "\n".join(parts)

    def is_empty(self) -> bool:
        return not (self.title or self.body_text or self.table_cells)


def extract_pptx(path: str | Path) -> list[RawSlide]:
    """
    Extract all slides from a .pptx file.
    Returns list of RawSlide objects ordered by slide number.
    """
    try:
        from pptx import Presentation
        from pptx.util import Pt
    except ImportError:
        raise RuntimeError("python-pptx not installed. Run: pip install python-pptx")

    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"PPTX not found: {path}")

    prs = Presentation(str(path))
    slides: list[RawSlide] = []

    for i, slide in enumerate(prs.slides, start=1):
        raw = RawSlide(slide_number=i)

        # Title
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            raw.title = slide.shapes.title.text_frame.text.strip()

        # All text frames
        body_parts = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = para.text.strip()
                    if line:
                        body_parts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        txt = cell.text.strip()
                        if txt:
                            raw.table_cells.append(txt)

        raw.body_text = "\n".join(body_parts)

        # Speaker notes
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            if notes_tf:
                raw.speaker_notes = notes_tf.text.strip()

        slides.append(raw)
        logger.debug(f"Slide {i}: '{raw.title}' — {len(body_parts)} text blocks")

    logger.info(f"Extracted {len(slides)} slides from {path.name}")
    return slides
